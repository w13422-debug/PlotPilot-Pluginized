"""Pure export logic adapted from the accepted PlotPilot v4.6.0 baseline.

The caller must supply immutable current-revision values obtained from Core.
This module never owns, saves, or mutates chapter text and returns bytes that
the host can publish through ``host.asset.create/v1``.
"""
from __future__ import annotations

import html
import io
import os
import re
import tempfile
from dataclasses import dataclass
from enum import Enum
from hashlib import sha256
from pathlib import Path
from typing import Iterable


class ExportFormat(str, Enum):
    EPUB = "epub"
    PDF = "pdf"
    DOCX = "docx"
    MARKDOWN = "markdown"


@dataclass(frozen=True, slots=True)
class ChapterRevision:
    document_id: str
    revision_id: str
    number: int
    title: str
    content: str
    content_hash: str

    def __post_init__(self) -> None:
        if self.number < 0:
            raise ValueError("chapter number must be non-negative")
        if sha256(self.content.encode("utf-8")).hexdigest() != self.content_hash:
            raise ValueError("chapter content does not match current revision hash")


@dataclass(frozen=True, slots=True)
class ExportDocument:
    workspace_id: str
    novel_id: str
    title: str
    author: str
    premise: str
    chapters: tuple[ChapterRevision, ...]


@dataclass(frozen=True, slots=True)
class ExportSelection:
    legacy_enabled: bool
    plugin_enabled: bool


@dataclass(frozen=True, slots=True)
class ExportPayload:
    content: bytes
    media_type: str
    filename: str
    sha256: str
    source_revisions: tuple[tuple[str, str, str], ...]


def validate_selection(selection: ExportSelection) -> str:
    """Enforce the frozen single-writer capability flag."""
    if selection.legacy_enabled == selection.plugin_enabled:
        raise ValueError("exactly one of legacy or plugin exporter must be enabled")
    return "plugin" if selection.plugin_enabled else "legacy"


def safe_filename_stem(title: str, max_len: int = 80) -> str:
    value = (title or "novel").strip()
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value)
    value = value.replace(" ", "_").strip("._") or "novel"
    return value[:max_len]


def chapter_display_title(chapter: ChapterRevision) -> str:
    return chapter.title.strip() if chapter.title and chapter.title.strip() else f"第 {chapter.number} 章"


def _ordered(chapters: Iterable[ChapterRevision]) -> tuple[ChapterRevision, ...]:
    # Stable sorting retains Core order for equal chapter numbers.
    return tuple(sorted(chapters, key=lambda chapter: chapter.number))


def _markdown(document: ExportDocument, chapters: tuple[ChapterRevision, ...]) -> tuple[bytes, str]:
    lines = [
        f"# {document.title or '未命名'}",
        "",
        f"**作者**: {document.author or '—'}",
        "",
        "## 简介",
        "",
        document.premise.strip() or "（无）",
        "",
    ]
    for chapter in chapters:
        lines.extend((f"## {chapter_display_title(chapter)}", "", chapter.content.strip() or "（无正文）", ""))
    return "\n".join(lines).encode("utf-8"), "text/markdown; charset=utf-8"


def _docx(document: ExportDocument, chapters: tuple[ChapterRevision, ...]) -> tuple[bytes, str]:
    from docx import Document

    output = Document()
    output.add_heading(document.title or "未命名", level=0)
    output.add_paragraph(f"作者：{document.author or '—'}")
    premise = output.add_paragraph()
    premise.add_run("简介：").bold = True
    premise.add_run(document.premise.strip() or "（无）")
    for chapter in chapters:
        output.add_heading(chapter_display_title(chapter), level=1)
        if not chapter.content.strip():
            output.add_paragraph("（无正文）")
        else:
            for line in chapter.content.splitlines():
                output.add_paragraph(line)
    stream = io.BytesIO()
    output.save(stream)
    return stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _epub(document: ExportDocument, chapters: tuple[ChapterRevision, ...]) -> tuple[bytes, str]:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier(f"plotpilot:{document.novel_id}")
    book.set_title(document.title or "未命名")
    book.set_language("zh")
    book.add_author(document.author or "未知作者")
    intro = epub.EpubHtml(title="简介", file_name="intro.xhtml", lang="zh")
    intro.content = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" lang="zh">'
        f"<head><title>简介</title><meta charset=\"utf-8\"/></head><body><h1>{html.escape(document.title or '未命名')}</h1>"
        f"<p>作者：{html.escape(document.author or '—')}</p><p>{html.escape(document.premise.strip() or '（无简介）')}</p></body></html>"
    )
    book.add_item(intro)
    spine = [intro]
    for index, chapter in enumerate(chapters, 1):
        title = chapter_display_title(chapter)
        paragraphs = [f"<p>{html.escape(line.strip())}</p>" for line in chapter.content.replace("\r\n", "\n").replace("\r", "\n").split("\n") if line.strip()]
        item = epub.EpubHtml(title=title, file_name=f"chap_{index:03d}.xhtml", lang="zh")
        item.content = (
            '<?xml version="1.0" encoding="utf-8"?>\n'
            '<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="zh">'
            f"<head><title>{html.escape(title)}</title><meta charset=\"utf-8\"/></head><body><h1>{html.escape(title)}</h1>"
            f"{'\n'.join(paragraphs) if paragraphs else '<p></p>'}</body></html>"
        )
        book.add_item(item)
        spine.append(item)
    book.toc = tuple(spine)
    book.add_item(epub.EpubNcx())
    book.spine = spine
    path: str | None = None
    try:
        descriptor, path = tempfile.mkstemp(suffix=".epub")
        os.close(descriptor)
        epub.write_epub(path, book, {})
        with open(path, "rb") as stream:
            content = stream.read()
    finally:
        if path and os.path.isfile(path):
            os.unlink(path)
    return content, "application/epub+zip"


def _pdf(document: ExportDocument, chapters: tuple[ChapterRevision, ...]) -> tuple[bytes, str]:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=14)
    font = ""
    candidates: list[Path] = []
    configured = (os.getenv("PLOTPILOT_EXPORT_CJK_FONT", "") or "").strip()
    if configured:
        candidates.append(Path(configured))
    if os.name == "nt":
        fonts = Path(os.getenv("WINDIR", r"C:\Windows")) / "Fonts"
        candidates.extend(fonts / name for name in ("msyh.ttf", "simhei.ttf", "simsun.ttc", "msyh.ttc", "simkai.ttf"))
    else:
        candidates.extend(Path(value) for value in (
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttf",
            "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        ))
    for candidate in candidates:
        if not candidate.is_file():
            continue
        try:
            pdf.add_font("PlotExportCJK", "", str(candidate), uni=True)
            font = "PlotExportCJK"
            break
        except Exception:
            continue
    if not font:
        raise RuntimeError("no usable CJK font is available for PDF export")
    pdf.add_page()

    def add(size: float, text: str, height: float) -> None:
        pdf.set_font(font, size=size)
        pdf.multi_cell(0, height, text or " ", new_x="LMARGIN", new_y="NEXT")

    add(16, document.title or "未命名", 9)
    add(11, f"作者：{document.author or '—'}\n简介：{document.premise.strip() or '—'}", 6)
    for chapter in chapters:
        add(14, chapter_display_title(chapter), 8)
        add(11, chapter.content.strip() or "（无正文）", 6)
    raw = pdf.output()
    return bytes(raw) if not isinstance(raw, str) else raw.encode("latin-1"), "application/pdf"


_RENDERERS = {
    ExportFormat.MARKDOWN: (_markdown, "md"),
    ExportFormat.DOCX: (_docx, "docx"),
    ExportFormat.EPUB: (_epub, "epub"),
    ExportFormat.PDF: (_pdf, "pdf"),
}


def build_export(document: ExportDocument, export_format: ExportFormat, *, document_id: str | None = None) -> ExportPayload:
    chapters = _ordered(document.chapters)
    if document_id is not None:
        chapters = tuple(chapter for chapter in chapters if chapter.document_id == document_id)
        if not chapters:
            raise ValueError(f"chapter does not exist: {document_id}")
        if len(chapters) != 1:
            raise ValueError(f"document id is ambiguous: {document_id}")
    renderer, extension = _RENDERERS[export_format]
    content, media_type = renderer(document, chapters)
    stem = safe_filename_stem(document.title)
    if document_id is not None:
        stem = safe_filename_stem(f"{document.title or 'novel'}-第{chapters[0].number}章")
    sources = tuple((chapter.document_id, chapter.revision_id, chapter.content_hash) for chapter in chapters)
    return ExportPayload(content, media_type, f"{stem}.{extension}", sha256(content).hexdigest(), sources)
